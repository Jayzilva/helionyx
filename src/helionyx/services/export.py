"""Exports: HOMER-importable series, Markdown reports, time series and scenario YAML (FR-EXP, FR-RPT-006)."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import jinja2
import numpy as np
import pandas as pd

from helionyx import DISCLAIMER
from helionyx.core.engine import dispatch as dsp
from helionyx.errors import validation
from helionyx.infra.db import now_iso
from helionyx.services.context import Helionyx
from helionyx.services.results import explain_run, get_monthly_summary, get_results
from helionyx.services.run import build_inputs, get_run
from helionyx.services.scenario import get_scenario, resolved_of, scenario_to_yaml

TEMPLATES = Path(__file__).resolve().parent.parent / "templates"
_env = jinja2.Environment(loader=jinja2.FileSystemLoader(TEMPLATES), autoescape=False, trim_blocks=True,
                          lstrip_blocks=True, undefined=jinja2.StrictUndefined)


def _fmt(v: Any, digits: int = 0) -> str:
    if v is None:
        return "—"
    if isinstance(v, float | int):
        return f"{v:,.{digits}f}"
    return str(v)


_env.filters["n0"] = lambda v: _fmt(v, 0)
_env.filters["n1"] = lambda v: _fmt(v, 1)
_env.filters["n2"] = lambda v: _fmt(v, 2)


def _out_dir(app: Helionyx, *parts: str) -> Path:
    d = app.settings.export_dir.joinpath(*parts)
    d.mkdir(parents=True, exist_ok=True)
    return d


def export_homer_csv(app: Helionyx, scenario_id: str | None = None, run_id: str | None = None) -> dict[str, Any]:
    if (scenario_id is None) == (run_id is None):
        raise validation("Give exactly one of scenario_id or run_id.", "Use the scenario you want to cross-check.")
    if run_id:
        scenario_id = get_run(app, run_id)["scenario_id"]
    assert scenario_id is not None
    doc = get_scenario(app, scenario_id)
    s = resolved_of(doc)
    inp = build_inputs(app, s)
    from helionyx.services.resource import dataset_frame

    res = dataset_frame(app, s.resource.dataset_id)
    out = _out_dir(app, scenario_id, "homer")
    files = {
        "load_kw.txt": inp.load,
        "ghi_kw_m2.txt": res["ghi_w_m2"].to_numpy() / 1000.0,
        "temp_c.txt": res["temp_c"].to_numpy(),
    }
    wind_height = None
    if "wind50_m_s" in res and s.components.wind is not None \
            and abs(s.components.wind.technical["hub_height_m"] - 50) < abs(s.components.wind.technical["hub_height_m"]
                                                                            - 10):
        files["wind_m_s.txt"] = res["wind50_m_s"].to_numpy()
        wind_height = 50
    else:
        files["wind_m_s.txt"] = res["wind10_m_s"].to_numpy()
        wind_height = 10
    for name, arr in files.items():
        np.savetxt(out / name, np.asarray(arr, dtype=np.float64), fmt="%.4f")
    params = _env.get_template("homer_parameters.md.j2").render(
        s=s.model_dump(mode="json"), scenario_id=scenario_id, scenario_hash=doc["scenario_hash"],
        wind_height=wind_height, generated_at=now_iso(), disclaimer=DISCLAIMER, assumptions=doc["assumptions"])
    (out / "homer_parameters.md").write_text(params, encoding="utf-8")
    totals = {"load_kwh_per_yr": float(inp.load.sum()), "ghi_kwh_m2_per_yr": float(files["ghi_kw_m2.txt"].sum()),
              "mean_temp_c": float(np.mean(files["temp_c.txt"])),
              "mean_wind_m_s": float(np.mean(files["wind_m_s.txt"]))}
    return {"scenario_id": scenario_id, "directory": str(out),
            "files": [str(out / n) for n in [*files, "homer_parameters.md"]],
            "anemometer_height_m": wind_height, "annual_totals": totals,
            "notes": ["Each series file has 8,760 rows starting 1 January 00:00 local time.",
                      "Import the series in HOMER Pro and enter the values from homer_parameters.md by hand.",
                      "Format follows the HOMER Pro manual (one value per line; 8,760 rows from midnight 1 January; "
                      "GHI as average kW/m² per step). An import test in HOMER Pro itself is part of the v0.2 "
                      "parity study."]}


def export_timeseries(app: Helionyx, run_id: str, rank: int = 1, fmt: str = "csv") -> dict[str, Any]:
    from helionyx.services.results import candidate_series

    c, series, _ = candidate_series(app, run_id, rank)
    df = pd.DataFrame(series, columns=list(dsp.SERIES_FIELDS))
    out = _out_dir(app, run_id)
    if fmt == "csv":
        path = out / f"rank{rank}_timeseries.csv"
        df.to_csv(path, index_label="hour")
    elif fmt == "parquet":
        path = out / f"rank{rank}_timeseries.parquet"
        df.to_parquet(path)
    else:
        raise validation("format must be csv or parquet.", "Use csv or parquet.")
    return {"run_id": run_id, "rank": rank, "path": str(path), "rows": len(df), "columns": list(df.columns)}


def report_data(app: Helionyx, run_id: str) -> dict[str, Any]:
    """Everything a report contains (FR-RPT-006), shared by the Markdown and Excel renderers."""
    from helionyx.services.results import _uncertainties

    run = get_run(app, run_id)
    scn = get_scenario(app, run["scenario_id"])
    batches = [app.db.get("batches", r["id"], "batch") for r in
               app.db.query("SELECT id FROM batches WHERE scenario_id = ? ORDER BY created_at", (run["scenario_id"],))]
    return {
        "run": run, "scn": scn, "s": scn["resolved"],
        "results": get_results(app, run_id, top_n=5),
        "ex": explain_run(app, run_id, rank=1, compare_to="base" if run.get("base_case") else "next"),
        "monthly": get_monthly_summary(app, run_id, rank=1),
        "sensitivities": [b["result"] for b in batches if b.get("status") == "completed"],
        "uncertainties": _uncertainties(app, scn),
        "generated_at": now_iso(), "disclaimer": DISCLAIMER,
        "homer_differences": (TEMPLATES / "differences_from_homer.md").read_text(encoding="utf-8"),
    }


def render_report(app: Helionyx, run_id: str) -> str:
    return _env.get_template("report.md.j2").render(**report_data(app, run_id))


def _cell(v: Any) -> Any:
    if v is None or isinstance(v, int | float | str):
        return v
    return json.dumps(v, default=str)


def write_xlsx(data: dict[str, Any], path: Path) -> None:
    """Excel report with one sheet per report section (FR-RPT-006, v0.2)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    run, scn, s, res, ex, monthly = (data[k] for k in ("run", "scn", "s", "results", "ex", "monthly"))
    cur = run["currency"]
    wb = Workbook()
    bold = Font(bold=True)
    first = [True]

    def sheet(title: str, header: list[str], rows: list[list[Any]]) -> None:
        ws = wb.active if first[0] else wb.create_sheet()
        first[0] = False
        assert ws is not None
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        for r in rows:
            ws.append([_cell(v) for v in r])
        for k, h in enumerate(header, start=1):
            width = max([len(str(h))] + [len(str(r[k - 1])) for r in rows if k - 1 < len(r)])
            ws.column_dimensions[get_column_letter(k)].width = min(60, max(10, width + 2))
        ws.freeze_panes = "A2"

    grid = s["grid"]
    tariff_txt = f", tariff {grid['tariff']['id']}, {grid['export_scheme']}" if grid["tariff"] else ""
    site = s["site"]
    sheet("Summary", ["Item", "Value"], [
        ["Report", f"Helionyx pre-feasibility report: {s['name']}"],
        ["Run", run["run_id"]], ["Scenario", f"{run['scenario_id']} (version {scn['version']})"],
        ["Scenario hash", run["scenario_hash"]], ["Engine", run["engine"]],
        ["Data packs", ", ".join(f"{k}@{v}" for k, v in run["packs"].items())],
        ["Generated", data["generated_at"]],
        ["Site", f"{site['name']} ({site['latitude']}, {site['longitude']}), {site['timezone']}"],
        ["Grid", grid["mode"] + tariff_txt], ["Grid availability", grid["availability"]["type"]],
        ["Dispatch", s["dispatch"]["strategy"]],
        ["Project life (years)", s["economics"]["project_life_years"]],
        ["Real discount rate", s["economics"]["real_discount_rate"]],
        ["Candidates", run["feasible_count"] + run["infeasible_count"]], ["Feasible", run["feasible_count"]],
        ["Currency", cur], ["Disclaimer", data["disclaimer"]],
    ])
    sheet("Assumptions", ["Field", "Value", "Origin", "Source", "Date"],
          [[a["path"], a["value"], a["origin"], a.get("source"), a.get("date")] for a in scn["assumptions"]]
          + [[f"warning {w['code']}", w["message"], "", "", ""] for w in scn["warnings"]])
    size_keys = ("pv_kwp", "wind_count", "bess_kwh", "bess_kw", "genset_kw")
    cols = ["npc", "lcoe_per_kwh", "initial_capital", "operating_cost_per_yr", "renewable_fraction_pct",
            "capacity_shortage_pct", "annual_bill", "simple_payback_yr", "irr_pct", "fuel_l_per_yr", "co2_kg_per_yr"]
    rows = [[c["rank"], *[c["sizes"][k] for k in size_keys], *[c["metrics"].get(k) for k in cols]]
            for c in res["candidates"]]
    if res["base_case"]:
        b = res["base_case"]
        rows.append([f"base: {b['label']}", *[b["sizes"][k] for k in size_keys], *[b["metrics"].get(k) for k in cols]])
    money = {"npc", "initial_capital", "operating_cost_per_yr", "annual_bill"}
    sheet("Top designs", ["Rank", "PV kWp", "Wind turbines", "BESS kWh", "BESS kW", "Genset kW",
                          *[f"{k} ({cur})" if k in money else k for k in cols]], rows)
    sheet("Cost breakdown", ["Scope", "Item", f"Present value ({cur})", "Share of gross cost (%)"],
          [["cost type", k, v["value"], v["share_pct"]] for k, v in ex["npc_by_cost_type"].items()]
          + [["component", k, v, None] for k, v in ex["npc_by_component"].items()]
          + [["total", "NPC", ex["npc"], None]])
    mkeys = [k for k in monthly["months"][0] if k != "month"]
    sheet("Monthly (rank 1)", ["Month", *mkeys],
          [[m["month"], *[m[k] for k in mkeys]] for m in monthly["months"]]
          + [["Total", *[monthly["annual_totals"][k] for k in mkeys]]])
    sens_rows: list[list[Any]] = []
    for b in data["sensitivities"]:
        for c in b["cases"]:
            vals = c.get("values") or {b.get("variable"): c.get("value")}
            sz, mt = c.get("optimal_sizes") or {}, c.get("optimal_metrics") or {}
            sens_rows.append([b["batch_id"], vals, c.get("optimal_architecture"), sz.get("pv_kwp"),
                              sz.get("bess_kwh"), sz.get("genset_kw"), mt.get("npc"), mt.get("lcoe_per_kwh")])
        for e in b["elasticities"]:
            sens_rows.append([b["batch_id"], f"NPC elasticity to {e['path']}", None, None, None, None,
                              e["npc_elasticity"], None])
    sheet("Sensitivity", ["Batch", "Values", "Architecture", "PV kWp", "BESS kWh", "Genset kW", f"NPC ({cur})",
                          "LCOE"], sens_rows or [["none", "No sensitivity analysis run", None, None, None, None,
                                                  None, None]])
    prov = res["provenance"]
    sheet("Provenance", ["Item", "Value"],
          [["Engine", prov["engine"]], ["Solver", prov["solver"]], ["Seed", prov["seed"]],
           ["Sources", ", ".join(prov["sources"])], ["Wall time (s)", prov.get("wall_time_s")],
           ["Uncertain inputs", "; ".join(f"{u['input']}: {u['reason']}" for u in data["uncertainties"])],
           ["Differences from HOMER", data["homer_differences"]], ["Disclaimer", data["disclaimer"]]])
    wb.save(path)


def export_report(app: Helionyx, run_id: str, fmt: str = "md") -> dict[str, Any]:
    if fmt not in ("md", "xlsx"):
        raise validation("format must be 'md' or 'xlsx'.", "Use md or xlsx.")
    out = _out_dir(app, run_id)
    if fmt == "md":
        text = render_report(app, run_id)
        path = out / "report.md"
        path.write_text(text, encoding="utf-8")
        return {"run_id": run_id, "format": "md", "path": str(path),
                "resource_uri": f"hnx://runs/{run_id}/report.md", "bytes": len(text.encode("utf-8"))}
    path = out / "report.xlsx"
    write_xlsx(report_data(app, run_id), path)
    return {"run_id": run_id, "format": "xlsx", "path": str(path), "resource_uri": None,
            "bytes": path.stat().st_size}


def export_scenario_yaml(app: Helionyx, scenario_id: str) -> str:
    return scenario_to_yaml(get_scenario(app, scenario_id))


def results_csv(app: Helionyx, run_id: str) -> str:
    get_run(app, run_id)
    rows = []
    for c in app.db.get_candidates(run_id, order="index"):
        row = {"candidate_index": c["candidate_index"], "rank": c.get("rank"), "feasible": c["feasible"],
               **c["sizes"], **c["metrics"],
               "violations": ";".join(v["constraint"] for v in c["violations"])}
        rows.append(row)
    return str(pd.DataFrame(rows).to_csv(index=False))
