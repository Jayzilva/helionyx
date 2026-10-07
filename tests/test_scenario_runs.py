"""Scenarios, optimiser, jobs, sensitivity, results and exports (FR-SCN, FR-OPT, FR-JOB, FR-SEN, FR-RPT, FR-EXP)."""

from __future__ import annotations

import time
from pathlib import Path

import pytest

from helionyx.core.models.scenario import ScenarioInput
from helionyx.errors import HelionyxError
from helionyx.infra.jobs import JobManager
from helionyx.services import export, results, run, scenario, sensitivity
from tests.conftest import create, hotel_input, offgrid_input, setup_site


def run_to_end(app, scenario_id, **kw):
    job = run.start_run(app, scenario_id, **kw)
    st = app.jobs.wait(job["job_id"], 300)
    assert st["state"] == "completed", st
    return job["run_id"]


# ------------------------------------------------------------------ scenario


def test_assumption_audit_lists_every_default(shared_app, hotel_ids):
    minimal = ScenarioInput.model_validate({
        "name": "minimal", "site_id": hotel_ids["site_id"], "load_ids": [hotel_ids["load_id"]],
        "resource_id": hotel_ids["resource_id"], "grid": {"mode": "grid_connected", "tariff_id": "lk.ceb.H2"},
        "components": {"pv": {"sizes_kwp": [0, 100]}}})
    doc = create(shared_app, minimal)
    v = scenario.validate_scenario(shared_app, doc["scenario_id"])
    paths = {a["path"]: a for a in v["assumptions"]}
    for p in ("grid.export_scheme", "economics.nominal_discount_rate", "economics.project_life_years",
              "components.pv.capital_per_unit", "components.pv.tilt_deg", "dispatch.strategy",
              "constraints.max_capacity_shortage", "seed", "economics.grid_kg_co2_per_kwh"):
        assert p in paths, p
        assert paths[p]["origin"] in ("pack_default", "system_default")
        assert paths[p]["source"] and paths[p]["date"]
    assert any(w["code"] == "HNX-W001" for w in v["warnings"])  # unverified tariff
    assert any(w["code"] == "HNX-W002" for w in v["warnings"])  # synthetic load


def test_override_recorded_as_user(shared_app, hotel_ids):
    inp = hotel_input(hotel_ids, components={"pv": {"sizes_kwp": [0, 100], "overrides": {"capital_per_unit": 1.0}}})
    doc = create(shared_app, inp)
    a = {x["path"]: x for x in doc["assumptions"]}
    assert a["components.pv.capital_per_unit"]["origin"] == "user"
    assert doc["resolved"]["components"]["pv"]["capital_per_unit"] == 1.0


@pytest.mark.parametrize("over, code", [
    ({"grid": {"mode": "grid_connected", "export_scheme": "net_plus", "tariff_id": "lk.ceb.H2"}}, "HNX-E008"),
    ({"grid": {"mode": "grid_connected"}}, "HNX-E001"),
    ({"components": {"bess": {"sizes_kwh": [100], "power_kw_per_kwh": 2.0}, "pv": {"sizes_kwp": [0]}}}, "HNX-E001"),
    ({"components": {"bess": {"sizes_kwh": [100], "overrides": {"round_trip_efficiency": 0.3}}}}, "HNX-E001"),
])
def test_validation_rules(shared_app, hotel_ids, over, code):
    doc = create(shared_app, hotel_input(hotel_ids, **over))
    assert any(e["code"] == code for e in doc["errors"]), doc["errors"]
    with pytest.raises(HelionyxError):
        run.start_run(shared_app, doc["scenario_id"])


def test_offgrid_without_storage_warns(shared_app, village_ids):
    doc = create(shared_app, offgrid_input(village_ids, components={"pv": {"sizes_kwp": [0, 50]}}))
    assert any(w["code"] == "HNX-W004" for w in doc["warnings"])


def test_search_space_limit(shared_app, hotel_ids):
    big = hotel_input(hotel_ids, components={"pv": {"sizes_kwp": {"min": 0, "max": 400, "step": 1}},
                                             "bess": {"sizes_kwh": {"min": 0, "max": 400, "step": 1}},
                                             "genset": {"sizes_kw": {"min": 0, "max": 400, "step": 1}}})
    with pytest.raises(HelionyxError) as e:
        create(shared_app, big)
    assert e.value.code.value == "HNX-E004" and e.value.details["candidate_count"] == 401 ** 3


def test_hash_stable_and_yaml_round_trip(shared_app, hotel_ids):
    a = create(shared_app, hotel_input(hotel_ids))
    b = create(shared_app, hotel_input(hotel_ids))
    assert a["scenario_hash"] == b["scenario_hash"] and a["candidate_count"] == 9
    text = scenario.scenario_to_yaml(a)
    c = create(shared_app, scenario.scenario_from_yaml(text))
    assert c["scenario_hash"] == a["scenario_hash"]


def test_versioning_after_run(shared_app, hotel_ids):
    a = create(shared_app, hotel_input(hotel_ids, components={"pv": {"sizes_kwp": [0, 50]}}))
    run_to_end(shared_app, a["scenario_id"])
    assert scenario.get_scenario(shared_app, a["scenario_id"])["locked"]
    b = scenario.create_scenario(shared_app, hotel_input(hotel_ids, components={"pv": {"sizes_kwp": [0, 60]}}),
                                 parent_scenario_id=a["scenario_id"])
    assert b["version"] == 2 and b["root_id"] == a["root_id"] and b["scenario_hash"] != a["scenario_hash"]


# ------------------------------------------------------------------ optimiser


def test_ranking_matches_brute_force(shared_app, hotel_ids):
    doc = create(shared_app, hotel_input(hotel_ids))
    run_id = run_to_end(shared_app, doc["scenario_id"])
    cands = shared_app.db.get_candidates(run_id, order="index")
    feas = sorted((c for c in cands if c["feasible"]), key=lambda c: (c["metrics"]["npc"], c["candidate_index"]))
    assert [c["candidate_index"] for c in feas] == \
        [c["candidate_index"] for c in shared_app.db.get_candidates(run_id, feasible_only=True)]
    res = results.get_results(shared_app, run_id)
    assert res["base_case"]["label"] == "grid_only"
    assert res["disclaimer"] and res["provenance"]["scenario_hash"] == doc["scenario_hash"]
    assert all(c["metrics"]["max_energy_balance_error_kwh"] <= 1e-3 for c in cands)
    # breakdown sums to NPC (FR-ECO-003)
    top = results.candidate_by_rank(shared_app, run_id, 1)
    assert sum(top["cost_by_type"].values()) == pytest.approx(top["metrics"]["npc"], rel=1e-4)


def test_no_feasible_returns_nearest(shared_app, hotel_ids):
    doc = create(shared_app, hotel_input(hotel_ids, constraints={"min_renewable_fraction": 0.99}))
    run_id = run_to_end(shared_app, doc["scenario_id"])
    with pytest.raises(HelionyxError) as e:
        results.get_results(shared_app, run_id)
    assert e.value.code.value == "HNX-E005"
    near = e.value.details["nearest_candidate"]
    assert near["violations"][0]["constraint"] == "min_renewable_fraction"


def test_offgrid_constraint_and_diesel_base(shared_app, village_ids):
    doc = create(shared_app, offgrid_input(village_ids))
    run_id = run_to_end(shared_app, doc["scenario_id"])
    res = results.get_results(shared_app, run_id, top_n=10)
    assert res["base_case"]["label"] == "diesel_only"
    for c in res["candidates"]:
        assert c["metrics"]["capacity_shortage_pct"] <= 1.0 + 1e-9
    ex = results.explain_run(shared_app, run_id, 1, "base")
    assert ex["comparison"]["against"] == "diesel_only"
    assert ex["dispatch_statistics"]["genset_hours_per_yr"] >= 0


def test_outage_reserve_increases_served_energy(shared_app, hotel_ids):
    def outage_run(reserve):
        inp = hotel_input(hotel_ids, components={"pv": {"sizes_kwp": [100]}, "bess": {"sizes_kwh": [200]}},
                          grid={"mode": "grid_connected", "tariff_id": "lk.ceb.H2",
                                "availability": {"type": "scheduled",
                                                 "windows": [{"days": "all", "start": "18:30", "end": "20:30"}]}},
                          dispatch={"grid": {"soc_reserve_for_outage": reserve}},
                          constraints={"max_capacity_shortage": 1.0})
        rid = run_to_end(shared_app, create(shared_app, inp)["scenario_id"])
        return results.candidate_by_rank(shared_app, rid, 1)["metrics"]
    low, high = outage_run(0.0), outage_run(0.6)
    assert low["outage_hours_per_yr"] == 2 * 365
    assert high["outage_unserved_kwh_per_yr"] < low["outage_unserved_kwh_per_yr"]


def test_monthly_summary_sums_to_annual(shared_app, hotel_ids):
    doc = create(shared_app, hotel_input(hotel_ids))
    run_id = run_to_end(shared_app, doc["scenario_id"])
    ms = results.get_monthly_summary(shared_app, run_id, 1)
    top = results.candidate_by_rank(shared_app, run_id, 1)
    assert ms["annual_totals"]["pv_kwh"] == pytest.approx(top["metrics"]["pv_kwh_per_yr"], rel=1e-9)
    assert ms["annual_totals"]["bill_after"] == pytest.approx(top["metrics"]["annual_bill"], rel=1e-9)
    base = shared_app.db.get("runs", run_id, "run")["base_case"]
    assert ms["annual_totals"]["bill_before"] == pytest.approx(base["metrics"]["annual_bill"], rel=1e-9)


def test_reports_and_homer_export(shared_app, hotel_ids):
    doc = create(shared_app, hotel_input(hotel_ids))
    run_id = run_to_end(shared_app, doc["scenario_id"])
    rep = export.export_report(shared_app, run_id, "md")
    text = Path(rep["path"]).read_text(encoding="utf-8")
    for section in ("Scenario summary", "Assumptions", "Top designs", "Cost breakdown", "Monthly energy",
                    "Sensitivity", "Provenance", "pre-feasibility"):
        assert section in text
    xl = export.export_report(shared_app, run_id, "xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(xl["path"])
    assert wb.sheetnames == ["Summary", "Assumptions", "Top designs", "Cost breakdown", "Monthly (rank 1)",
                             "Sensitivity", "Provenance"]
    assert any("pre-feasibility" in str(c.value) for row in wb["Summary"].iter_rows() for c in row)
    h = export.export_homer_csv(shared_app, run_id=run_id)
    lines = Path(h["files"][0]).read_text(encoding="utf-8").strip().splitlines()
    assert len(lines) == 8760
    assert h["annual_totals"]["load_kwh_per_yr"] == pytest.approx(540000, rel=1e-6)


def test_sensitivity_and_elasticity(shared_app, hotel_ids):
    doc = create(shared_app, hotel_input(hotel_ids))
    var = sensitivity.SweepVariable(path="components.pv.overrides.capital_per_unit", values=[100000, 250000])
    job = sensitivity.start_sensitivity(shared_app, doc["scenario_id"], [var])
    assert shared_app.jobs.wait(job["job_id"], 300)["state"] == "completed"
    out = sensitivity.get_sensitivity_results(shared_app, job["batch_id"])
    assert len(out["cases"]) == 2 and out["base_value"] == 160000
    el = out["elasticities"][0]["npc_elasticity"]
    assert el is not None and el > 0
    cheap, dear = out["cases"]
    assert cheap["optimal_sizes"]["pv_kwp"] >= dear["optimal_sizes"]["pv_kwp"]
    with pytest.raises(HelionyxError) as e:
        sensitivity.start_sensitivity(shared_app, doc["scenario_id"], [var, var, var])
    assert e.value.code.value == "HNX-E008"


def test_two_variable_grid(shared_app, hotel_ids):
    doc = create(shared_app, hotel_input(hotel_ids))
    pv = sensitivity.SweepVariable(path="components.pv.overrides.capital_per_unit", values=[100000, 300000])
    gp = sensitivity.SweepVariable(path="economics.grid_price_escalation_real", values=[0.0, 0.05])
    job = sensitivity.start_sensitivity(shared_app, doc["scenario_id"], [pv, gp])
    assert job["cases"] == 4
    assert shared_app.jobs.wait(job["job_id"], 300)["state"] == "completed"
    out = sensitivity.get_sensitivity_results(shared_app, job["batch_id"])
    assert out["kind"] == "grid" and len(out["cases"]) == 4 and len(out["elasticities"]) == 2
    assert all(c["optimal_architecture"] for c in out["cases"])
    cells = {(c["values"][pv.path], c["values"][gp.path]): c for c in out["cases"]}
    # cheaper PV never leads to less PV in the optimal design
    assert cells[(100000, 0.0)]["optimal_sizes"]["pv_kwp"] >= cells[(300000, 0.0)]["optimal_sizes"]["pv_kwp"]


def test_compare_runs(shared_app, hotel_ids):
    a = run_to_end(shared_app, create(shared_app, hotel_input(hotel_ids))["scenario_id"])
    b = run_to_end(shared_app, create(shared_app, hotel_input(
        hotel_ids, components={"pv": {"sizes_kwp": [0, 150]}}))["scenario_id"])
    cmp_ = results.compare_runs(shared_app, [a, b])
    assert "difference_vs_first" in cmp_["runs"][1]


# ------------------------------------------------------------------ jobs


def test_job_limit_cancel_and_interrupt(app):
    jm = JobManager(app.db, max_concurrent=1, timeout_s=60)

    def slow(progress):
        for i in range(200):
            progress(i / 2, "working")
            time.sleep(0.01)
        return {}

    j = jm.submit("test", "u", slow, {})
    with pytest.raises(HelionyxError) as e:
        jm.submit("test", "u", slow, {})
    assert e.value.code.value == "HNX-E007"
    jm.cancel(j["job_id"])
    assert jm.wait(j["job_id"], 10)["state"] == "cancelled"
    # a job left 'running' in the store is marked interrupted on restart
    doc = dict(jm.status(j["job_id"]), state="running")
    app.db.put("jobs", doc["job_id"], doc, kind="test", state="running", owner="u")
    jm2 = JobManager(app.db, 1, 60)
    assert jm2.status(j["job_id"])["state"] == "interrupted"
    jm.shutdown()
    jm2.shutdown()


def test_job_timeout(app):
    jm = JobManager(app.db, max_concurrent=2, timeout_s=0.05)

    def slow(progress):
        for _ in range(100):
            time.sleep(0.01)
            progress(1, "x")
        return {}

    j = jm.submit("test", "u", slow, {})
    st = jm.wait(j["job_id"], 10)
    assert st["state"] == "failed" and st["error"]["code"] == "HNX-E006"
    jm.shutdown()


def test_start_run_returns_quickly(app):
    ids = setup_site(app)
    doc = create(app, hotel_input(ids))
    t = time.perf_counter()
    job = run.start_run(app, doc["scenario_id"])
    assert time.perf_counter() - t < 2.0
    app.jobs.wait(job["job_id"], 300)


@pytest.mark.slow
def test_benchmark_1000_candidates(shared_app, village_ids):
    """NFR-PERF-01: a 1,000-candidate search over 8,760 hours in ≤ 30 s (warm)."""
    inp = offgrid_input(village_ids, components={
        "pv": {"sizes_kwp": {"min": 0, "max": 225, "step": 25}},
        "bess": {"sizes_kwh": {"min": 0, "max": 900, "step": 100}},
        "genset": {"sizes_kw": [0, 20, 30, 40, 50, 60, 70, 80, 90, 100]}})
    doc = create(shared_app, inp)
    assert doc["candidate_count"] == 1000
    run_id = run_to_end(shared_app, doc["scenario_id"])
    wall = shared_app.db.get("runs", run_id, "run")["wall_time_s"]
    print(f"1,000-candidate run: {wall:.2f} s")
    assert wall <= 30.0
